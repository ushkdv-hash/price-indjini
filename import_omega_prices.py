# -*- coding: utf-8 -*-
"""Импорт цен «Омега-Принт» из файла «Инджини цены для офсета (5).xlsx».

Создаёт новые калькуляторы:
  - listovki_omega          (Листовки, вкладка «Листовки», calc_type=offset_table)
  - buklety_omega           (Буклеты, вкладка «Буклеты», calc_type=offset_table)
  - bumazhnye_pakety_omega  (Бумажные пакеты, calc_type=paper_bags)

и переводит в статус disabled в data/offset_registry.json прежние офсет-позиции
(vizitki_ofset, listovki_ofset, buklety_liflety_ofset, papki_*, bumazhnye_pakety),
чтобы их можно было включить обратно в админке.

Запуск:
  python import_omega_prices.py            # сухой прогон (показать сводку)
  python import_omega_prices.py --apply    # записать products.json и offset_registry.json
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
from datetime import date, timedelta

try:
    import openpyxl
except ImportError:
    print("Установите openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(APP_DIR, "data")
PRODUCTS_PATH = os.path.join(DATA_DIR, "products.json")
REGISTRY_PATH = os.path.join(DATA_DIR, "offset_registry.json")

PRICE_FACTOR = 1.6
WORKING_DAYS = 7
SOURCE_NOTE = "Цены даны по прайсу типографии Омега-Принт"
PRICE_SOURCE = "Инджини цены для офсета (5).xlsx"

DISCOUNTS = [
    {"name": "Без скидки", "percent": 0},
    {"name": "Оптовый клиент", "percent": 20},
    {"name": "Дизайнер", "percent": 15},
    {"name": "Рекламное агентство", "percent": 15},
]

PAPER_BAGS_MANDATORY_OPTIONS = [
    {
        "id": "lamination",
        "label": "Ламинация",
        "choices": [
            {"id": "gloss", "name": "Глянцевая 1+0"},
            {"id": "matte", "name": "Матовая 1+0"},
        ],
    },
    {
        "id": "eyelets",
        "label": "Люверсы",
        "choices": [
            {"id": "gold", "name": "Золото"},
            {"id": "silver", "name": "Серебро"},
        ],
    },
    {
        "id": "cord",
        "label": "Шнурок",
        "choices": [
            {"id": "white", "name": "Белый"},
            {"id": "black", "name": "Чёрный"},
            {"id": "blue", "name": "Синий"},
            {"id": "red", "name": "Красный"},
        ],
    },
]

# Итоговый реестр офсет-позиций: новые включены, прежние отключены
# (включить обратно можно в админке: Админка → Офсет — прайсы).
REGISTRY_CATALOG = [
    {"id": "vizitki_ofset", "title": "Визитки", "tab": "Визитки", "status": "disabled"},
    {"id": "listovki_ofset", "title": "Листовки и плакаты", "tab": "Листовки и плакаты", "status": "disabled"},
    {"id": "listovki_omega", "title": "Листовки", "tab": "Листовки", "status": "active"},
    {"id": "buklety_liflety_ofset", "title": "Буклеты и лифлеты", "tab": "Буклеты и лифлеты", "status": "disabled"},
    {"id": "buklety_omega", "title": "Буклеты", "tab": "Буклеты", "status": "active"},
    {"id": "papki_0204_ofset", "title": "Папки №0204", "tab": "Папки", "status": "disabled"},
    {"id": "papki_145_ofset", "title": "Папка №145", "tab": "Папки", "status": "disabled"},
    {"id": "papki_835_ofset", "title": "Папка №835", "tab": "Папки", "status": "disabled"},
    {"id": "bumazhnye_pakety", "title": "Бумажные пакеты (прайс Сборка)", "tab": "Бумажные пакеты", "status": "disabled"},
    {"id": "bumazhnye_pakety_omega", "title": "Бумажные пакеты", "tab": "Бумажные пакеты", "status": "active"},
]

# Виды буклетов/лифлетов: значение в столбце «Вид» → (id, название)
BUKLETY_VARIANTS = {
    "Лифлет A3 (3 фальца)": ("liflet_a3_3fold", "Лифлет A3 (3 фальца)"),
    "Лифлет A3 в гармошку": ("liflet_a3_accordion", "Лифлет A3 в гармошку (2 фальца)"),
    "Лифлет A3 в накидку": ("liflet_a3_gate", "Лифлет A3 в накидку (2 фальца)"),
    "Лифлет A4 в гармошку (3 фальца)": ("liflet_a4_3fold", "Лифлет A4 в гармошку (3 фальца)"),
    "Лифлет A4 в гармошку": ("liflet_a4_accordion", "Лифлет A4 в гармошку (2 фальца)"),
    "Лифлет A4 в накидку": ("liflet_a4_gate", "Лифлет A4 в накидку (2 фальца)"),
    "Буклет A3 в A4": ("buklet_a3_a4", "Буклет A3 в A4 (1 фальц)"),
    "Буклет A4 в A5": ("buklet_a4_a5", "Буклет A4 в A5 (1 фальц)"),
    "Буклет A5 в A6": ("buklet_a5_a6", "Буклет A5 в A6 (1 фальц)"),
    "Буклет A6 в A7": ("buklet_a6_a7", "Буклет A6 в A7 (1 фальц)"),
    "Евробуклет": ("euro_buklet", "Евробуклет (210х198 в развороте, 1 фальц)"),
}

# Материалы: значение в столбце «Бумага» → (id, название)
MATERIALS = {
    "Мелованная 115 г/м²": ("115_mel", "Мелованная бумага 115 г/м²"),
    "Мелованная 130 г/м²": ("130_mel", "Бумага мелованная 130 г/м²"),
    "Мелованная 200 г/м²": ("200_mel", "Мелованная бумага 200 г/м²"),
    "Мелованный картон 300 г/м²": ("300_mel", "Мелованный картон 300 г/м²"),
    "Офсетная 80 г/м²": ("80_offset", "Офсетная 80 г/м²"),
}

LISTOVKI_SIZE_IDS = {
    "Евро (100×210 мм)": "euro",
    "A7 (74×105 мм)": "a7",
    "A6 (105×148 мм)": "a6",
    "A5 (148×210 мм)": "a5",
    "A4 (210×297 мм)": "a4",
    "A3 (297×420 мм)": "a3",
    "A2 (420×594 мм)": "a2",
}

QTY_LABEL_RE = re.compile(r"(\d[\d\s]*)\s*шт", re.IGNORECASE)

PAPER_BAG_CATEGORY = {
    "маленькие": "small",
    "средние": "medium",
    "большие": "large",
}

PAPER_BAG_CATEGORY_NAME = {
    "small": "маленькие",
    "medium": "средние",
    "large": "большие",
}

# Категории пакетов из шаблона прайса (tools/generate_contractor_template.py):
# «тройка» размеров в порядке листа (В×Ш×Г) → категория
PAPER_BAG_DIM_CATEGORY = {
    (148, 105, 50): "small",
    (240, 250, 75): "small",
    (320, 230, 80): "small",
    (190, 120, 90): "small",
    (320, 320, 100): "medium",
    (350, 250, 80): "medium",
    (360, 240, 100): "medium",
    (170, 230, 80): "medium",
    (350, 310, 190): "large",
    (450, 330, 160): "large",
    (355, 380, 100): "large",
    (235, 360, 100): "large",
    (400, 300, 120): "large",
}


def find_source_file():
    candidates = sorted(glob.glob(os.path.join(APP_DIR, "*офсета*.xlsx")))
    if not candidates:
        raise RuntimeError("Не найден файл «Инджини цены для офсета (5).xlsx» в папке проекта")
    return candidates[0]


def parse_qty_from_label(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    m = QTY_LABEL_RE.search(str(value).replace("\xa0", " "))
    if not m:
        return None
    digits = re.sub(r"[^\d]", "", m.group(1))
    return int(digits) if digits else None


def parse_price(value):
    """Цена за 1 шт из ячейки (None → цены нет)."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if float(value) > 0 else None
    s = str(value).replace("\u20bd", "").replace("₽", "").replace("\xa0", " ").strip()
    s = s.replace(" ", "").replace(",", ".")
    if not s or s in ("—", "-", "0", "0.0"):
        return None
    try:
        v = float(s)
    except ValueError:
        return None
    return v if v > 0 else None


def cell_str(value):
    return str(value).replace("\xa0", " ").strip() if value is not None else ""


def add_working_days(start, n):
    """Дата через n рабочих дней после start (пн–пт)."""
    d = start
    added = 0
    while added < n:
        d += timedelta(days=1)
        if d.weekday() < 5:
            added += 1
    return d


def price_as_of():
    return date.today()


def delivery_tiers():
    ref = price_as_of()
    end = add_working_days(ref, WORKING_DAYS - 1)  # 7 раб. дн. с учётом дня старта
    tiers = [
        {
            "id": "standard",
            "label": f"Готовность на {end.day:02d}.{end.month:02d}",
            "date_end": end.isoformat(),
        }
    ]
    return tiers, ref, end


def find_header(ws, marker, max_row=5):
    """Строка шапки (с marker) и словарь {индекс колонки: тираж}."""
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        cells = [cell_str(c.value) for c in row]
        if not any(marker in c for c in cells):
            continue
        qty_cols = {}
        qty_order = []
        for idx, c in enumerate(row):
            q = parse_qty_from_label(c.value)
            if q is not None:
                qty_cols[idx] = q
                if q not in qty_order:
                    qty_order.append(q)
        if qty_cols:
            return row[0].row, qty_cols, qty_order
        return row[0].row, {}, []
    return None, {}, []


def parse_offset_table_sheet(ws, variant_col=False):
    """Разбор листа-таблицы (Листовки / Буклеты и лифлеты).

    Возвращает (rows, qty_order);
    rows = [{format, variant, material, print_mode, unit_prices{qty: цена за шт}}].
    """
    header_row, qty_cols, qty_order = find_header(ws, "Цве")  # «Цветность» или «Цве-ть»
    if header_row is None or not qty_cols:
        raise RuntimeError(f"Не найдена шапка таблицы на листе «{ws.title}»")

    rows = []
    for row in ws.iter_rows(min_row=header_row + 1):
        fmt = cell_str(row[0].value)
        variant = cell_str(row[1].value) if variant_col else ""
        material = cell_str(row[2].value if variant_col else row[1].value)
        print_mode = cell_str(row[3].value if variant_col else row[2].value)
        if not fmt or not material or print_mode not in ("4+0", "4+4"):
            continue
        unit_prices = {}
        for idx, q in qty_cols.items():
            price = parse_price(row[idx].value)
            if price is not None:
                unit_prices[q] = price
        if not unit_prices:
            continue
        rows.append({
            "format": fmt,
            "variant": variant,
            "material": material,
            "print_mode": print_mode,
            "unit_prices": unit_prices,
        })
    return rows, qty_order


def mat_name(mat_id):
    return next(name for key, (mid, name) in MATERIALS.items() if mid == mat_id)


def build_listovki_product(rows, qty_order):
    tiers, ref, end = delivery_tiers()
    tier_id = tiers[0]["id"]

    sizes, size_ids = [], {}
    for r in rows:
        if r["format"] in size_ids:
            continue
        sid = LISTOVKI_SIZE_IDS.get(r["format"])
        if not sid:
            raise RuntimeError(f"Неизвестный формат листовки: {r['format']!r}")
        size_ids[r["format"]] = sid
        sizes.append({"id": sid, "name": r["format"]})

    seen_materials, prices = [], {}
    for r in rows:
        mat_id = MATERIALS[r["material"]][0]
        if mat_id not in seen_materials:
            seen_materials.append(mat_id)
        mode_block = (
            prices.setdefault(size_ids[r["format"]], {})
            .setdefault(mat_id, {})
            .setdefault(r["print_mode"], {})
        )
        for qty in qty_order:
            unit = r["unit_prices"].get(qty)
            if unit is None:
                continue
            mode_block.setdefault(str(qty), {})[tier_id] = round(unit * qty, 2)

    size_order = [s["id"] for s in sizes]
    materials = []
    for mid in seen_materials:
        used = sorted(
            {size_ids[r["format"]] for r in rows if MATERIALS[r["material"]][0] == mid},
            key=size_order.index,
        )
        materials.append({"id": mid, "name": mat_name(mid), "size_ids": used})

    return {
        "id": "listovki_omega",
        "title": "Листовки",
        "icon": "listovki",
        "calc_type": "offset_table",
        "offset_group": "ofset",
        "offset_tab": "Листовки",
        "price_factor": PRICE_FACTOR,
        "price_as_of": ref.isoformat(),
        "price_table_updated": ref.isoformat(),
        "price_source": PRICE_SOURCE,
        "price_source_note": SOURCE_NOTE,
        "delivery_tier_display": "working_days",
        "materials": materials,
        "print_modes": ["4+0", "4+4"],
        "print_mode_labels": {
            "4+0": "Односторонняя печать",
            "4+4": "Двусторонняя печать",
        },
        "qty_options": list(qty_order),
        "delivery_tiers": tiers,
        "sizes": sizes,
        "prices": prices,
        "options": [],
        "discounts": [dict(d) for d in DISCOUNTS],
    }


def build_buklety_product(rows, qty_order):
    tiers, ref, end = delivery_tiers()
    tier_id = tiers[0]["id"]

    sizes, size_ids, row_variants = [], {}, []
    for r in rows:
        variant = r["variant"]
        if variant not in BUKLETY_VARIANTS:
            raise RuntimeError(f"Неизвестный вид буклета: {variant!r}")
        vid, vname = BUKLETY_VARIANTS[variant]
        if vid not in size_ids:
            size_ids[vid] = vid
            sizes.append({"id": vid, "name": vname})
        row_variants.append(vid)

    seen_materials, prices = [], {}
    for r, vid in zip(rows, row_variants):
        mat_id = MATERIALS[r["material"]][0]
        if mat_id not in seen_materials:
            seen_materials.append(mat_id)
        mode_block = (
            prices.setdefault(vid, {})
            .setdefault(mat_id, {})
            .setdefault(r["print_mode"], {})
        )
        for qty in qty_order:
            unit = r["unit_prices"].get(qty)
            if unit is None:
                continue
            mode_block.setdefault(str(qty), {})[tier_id] = round(unit * qty, 2)

    materials = []
    for mid in seen_materials:
        used = list(dict.fromkeys(
            vid for r, vid in zip(rows, row_variants) if MATERIALS[r["material"]][0] == mid
        ))
        materials.append({"id": mid, "name": mat_name(mid), "size_ids": used})

    return {
        "id": "buklety_omega",
        "title": "Буклеты",
        "icon": "brosyury_menyu",
        "calc_type": "offset_table",
        "offset_group": "ofset",
        "offset_tab": "Буклеты",
        "size_label": "Вид",
        "price_factor": PRICE_FACTOR,
        "price_as_of": ref.isoformat(),
        "price_table_updated": ref.isoformat(),
        "price_source": PRICE_SOURCE,
        "price_source_note": SOURCE_NOTE,
        "delivery_tier_display": "working_days",
        "materials": materials,
        "print_modes": ["4+0", "4+4"],
        "print_mode_labels": {
            "4+0": "Односторонняя печать",
            "4+4": "Двусторонняя печать",
        },
        "qty_options": list(qty_order),
        "delivery_tiers": tiers,
        "sizes": sizes,
        "prices": prices,
        "options": [],
        "discounts": [dict(d) for d in DISCOUNTS],
    }




def parse_paper_bags_sheet(ws):
    """Разбор листа «Бумажные пакеты» → ([items], qty_order).

    Включает только размеры, для которых заполнены цены.
    """
    header_row, qty_cols, qty_order = find_header(ws, "Комплектация")
    if header_row is None or not qty_cols:
        raise RuntimeError(f"Не найдена шапка таблицы на листе «{ws.title}»")

    items = []
    for row in ws.iter_rows(min_row=header_row + 1):
        dims_text = cell_str(row[0].value)
        m = re.match(r"^(\d+)\s*[x×хX]\s*(\d+)\s*[x×хX]\s*(\d+)", dims_text)
        if not m:
            continue
        prices = {}
        for idx, q in qty_cols.items():
            unit = parse_price(row[idx].value)
            if unit is not None:
                prices[str(q)] = round(unit * q, 2)  # себестоимость за весь тираж
        if not prices:
            continue  # размер без цен не включаем
        h, w, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        low = dims_text.lower()
        category = next(
            (cat for key, cat in PAPER_BAG_CATEGORY.items() if key in low), None
        )
        if category is None:
            category = PAPER_BAG_DIM_CATEGORY.get((h, w, d), "other")
        prefix = {"small": "small", "medium": "medium", "large": "large"}.get(category, "bag")
        items.append({
            "id": f"{prefix}_{h}_{w}_{d}_мм",
            "name": f"Пакеты {PAPER_BAG_CATEGORY_NAME.get(category, dims_text)} ({h}x{w}x{d} мм)",
            "category": category,
            "qty_options": [q for q in qty_order if str(q) in prices],
            "prices": prices,
            "height_mm": h,
            "width_mm": w,
            "depth_mm": d,
        })
    return items, qty_order


def build_paper_bags_product(items):
    ref = price_as_of()
    end = add_working_days(ref, WORKING_DAYS - 1)
    sizes, prices = [], {}
    for item in items:
        sizes.append({
            "id": item["id"],
            "name": item["name"],
            "category": item["category"],
            "qty_options": item["qty_options"],
            "width_mm": item["width_mm"],
            "height_mm": item["height_mm"],
            "depth_mm": item["depth_mm"],
        })
        prices[item["id"]] = item["prices"]
    return {
        "id": "bumazhnye_pakety_omega",
        "title": "Бумажные пакеты",
        "icon": "bumazhnye_pakety",
        "calc_type": "paper_bags",
        "price_source": PRICE_SOURCE,
        "price_source_note": SOURCE_NOTE,
        "price_as_of": ref.isoformat(),
        "price_table_updated": ref.isoformat(),
        "delivery": {
            "label": f"Готовность на {end.day:02d}.{end.month:02d}",
            "date": end.isoformat(),
            "working_days": WORKING_DAYS,
            "display": f"{WORKING_DAYS} раб. дн.",
        },
        "sizes": sizes,
        "prices": prices,
        "mandatory_options": [dict(g) for g in PAPER_BAGS_MANDATORY_OPTIONS],
        "price_factor": PRICE_FACTOR,
        "options": [],
        "discounts": [dict(d) for d in DISCOUNTS],
    }


def load_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def save_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    arg_parser = argparse.ArgumentParser(description="Импорт цен Омега-Принт (Инджини)")
    arg_parser.add_argument("--apply", action="store_true", help="записать изменения в data/")
    args = arg_parser.parse_args()

    path = find_source_file()
    print("Источник:", os.path.basename(path))
    wb = openpyxl.load_workbook(path, data_only=True)

    rows_l, qty_l = parse_offset_table_sheet(wb["Листовки"])
    listovki = build_listovki_product(rows_l, qty_l)

    rows_b, qty_b = parse_offset_table_sheet(wb["Буклеты и лифлеты"], variant_col=True)
    buklety = build_buklety_product(rows_b, qty_b)

    bag_items, _bag_qty = parse_paper_bags_sheet(wb["Бумажные пакеты"])
    bags = build_paper_bags_product(bag_items)

    tiers, ref, end = delivery_tiers()
    wd = sum(
        1 for i in range((end - ref).days + 1)
        if (ref + timedelta(days=i)).weekday() < 5
    )
    assert wd == WORKING_DAYS, f"Срок считался как {wd} раб. дн."

    print(f"Листовки: {len(listovki['sizes'])} форматов, {len(listovki['materials'])} материала, "
          f"тиражи {listovki['qty_options']}")
    print(f"Буклеты:  {len(buklety['sizes'])} видов, {len(buklety['materials'])} материала, "
          f"тиражи {buklety['qty_options']}")
    print("Пакеты:   размеры с ценами:")
    for s in bags["sizes"]:
        print("   -", s["name"], "тиражи:", s["qty_options"])
    print(f"Срок производства: {WORKING_DAYS} раб. дн. "
          f"(прайс от {ref.isoformat()}, готовность {end.isoformat()})")

    if not args.apply:
        print("\nСухой прогон — данные не записаны. Для записи: python import_omega_prices.py --apply")
        return 0

    products = load_json(PRODUCTS_PATH)
    products["listovki_omega"] = listovki
    products["buklety_omega"] = buklety
    products["bumazhnye_pakety_omega"] = bags
    save_json(PRODUCTS_PATH, products)
    print("Записано:", PRODUCTS_PATH)

    registry = load_json(REGISTRY_PATH)
    registry["catalog"] = [dict(item) for item in REGISTRY_CATALOG]
    save_json(REGISTRY_PATH, registry)
    print("Записано:", REGISTRY_PATH)
    print("Отключены (можно включить в админке):",
          ", ".join(i["id"] for i in REGISTRY_CATALOG if i["status"] == "disabled"))
    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)

